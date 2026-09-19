# agents/multi_file_reader.py

import os
import json
import csv
from pathlib import Path
from typing import List, Dict, Union
import re

class MultiFileReaderAgent:
    """
    Agent de lecture multi-fichiers pour PME
    Supporte: PDF, DOCX, XLSX, CSV, TXT, JSON, PPTX, JPG, PNG, ZIP
    """
    
    def __init__(self, upload_folder="uploads", max_files=20, max_size_mb=50):
        self.upload_folder = upload_folder
        self.max_files = max_files
        self.max_size_mb = max_size_mb
        Path(upload_folder).mkdir(exist_ok=True)
        print(f"📁 MultiFileReaderAgent - Dossier: {upload_folder}")
        print(f"📋 Max fichiers: {max_files} | Max taille: {max_size_mb} Mo")
    
    # ============================================
    # LIRE UN SEUL FICHIER
    # ============================================
    def read_file(self, file_path: str) -> str:
        """Lit un seul fichier"""
        try:
            if not os.path.exists(file_path):
                return f"❌ Fichier introuvable: {file_path}"
            
            if not self._check_size(file_path):
                return f"⚠️ Fichier trop gros: {os.path.basename(file_path)}"

            ext = Path(file_path).suffix.lower()
            
            if ext == '.txt':
                return self.read_txt(file_path)
            elif ext == '.csv':
                return self.read_csv(file_path)
            elif ext == '.json':
                return self.read_json(file_path)
            elif ext in ['.pdf']:
                return self.read_pdf(file_path)
            elif ext in ['.docx', '.doc']:
                return self.read_docx(file_path)
            elif ext in ['.xlsx', '.xls']:
                return self.read_excel(file_path)
            elif ext in ['.pptx', '.ppt']:
                return self.read_pptx(file_path)
            elif ext in ['.jpg', '.jpeg', '.png', '.gif', '.bmp']:
                return self.read_image(file_path)
            elif ext in ['.zip', '.rar', '.7z']:
                return self.read_archive(file_path)
            else:
                return f"⚠️ Format non supporté: {ext}"
        except Exception as e:
            return f"❌ Erreur: {str(e)}"
    
    # ============================================
    # LIRE PLUSIEURS FICHIERS
    # ============================================
    def read_files(self, file_paths: Union[str, List[str]], combine: bool = True) -> Union[str, Dict]:
        """Lit plusieurs fichiers"""
        if isinstance(file_paths, str):
            file_paths = [file_paths]
        
        if len(file_paths) > self.max_files:
            return {
                'error': f"❌ Trop de fichiers (max {self.max_files})",
                'count': len(file_paths)
            }
        
        results = {}
        errors = []
        total_size = 0

        for i, path in enumerate(file_paths):
            if not os.path.exists(path):
                errors.append(f"⚠️ Fichier {i+1} introuvable: {path}")
                continue

            size_mb = os.path.getsize(path) / (1024 * 1024)
            total_size += size_mb

            if size_mb > self.max_size_mb:
                errors.append(f"⚠️ Fichier {i+1} trop gros: {os.path.basename(path)} ({size_mb:.1f} Mo)")
                continue

            content = self.read_file(path)
            filename = os.path.basename(path)
            results[filename] = content

        if total_size > self.max_size_mb * 2:
            errors.append(f"⚠️ Taille totale trop grande: {total_size:.1f} Mo")

        if combine:
            return self._combine_results(results, errors)
        else:
            return {
                'results': results,
                'errors': errors,
                'count': len(results),
                'total_size_mb': round(total_size, 2)
            }
    
    # ============================================
    # COMBINER LES RÉSULTATS
    # ============================================
    def _combine_results(self, results: Dict, errors: List) -> str:
        """Combine les résultats en un seul texte"""
        if not results and not errors:
            return "📭 Aucun fichier à lire"

        output = []
        output.append(f"📊 RÉSULTAT DE LECTURE ({len(results)} fichiers)")
        output.append("=" * 50)

        for filename, content in results.items():
            output.append(f"\n📄 {filename}")
            output.append("-" * 30)
            output.append(content)

        if errors:
            output.append("\n⚠️ ERREURS:")
            for err in errors:
                output.append(f"  {err}")

        return "\n".join(output)
    
    # ============================================
    # LIRE UN DOSSIER ENTIER
    # ============================================
    def read_folder(self, folder_path: str, extensions: List[str] = None) -> str:
        """Lit tous les fichiers d'un dossier"""
        if not os.path.exists(folder_path):
            return f"❌ Dossier introuvable: {folder_path}"

        files = []
        for root, dirs, filenames in os.walk(folder_path):
            for filename in filenames:
                ext = Path(filename).suffix.lower()
                if extensions is None or ext in extensions:
                    files.append(os.path.join(root, filename))

        if not files:
            return f"📭 Aucun fichier trouvé dans {folder_path}"

        return self.read_files(files)

    # ============================================
    # VÉRIFIER LA TAILLE
    # ============================================
    def _check_size(self, file_path: str) -> bool:
        try:
            size_mb = os.path.getsize(file_path) / (1024 * 1024)
            return size_mb <= self.max_size_mb
        except:
            return False

    # ============================================
    # MÉTHODE PROCESS (pour compatibilité avec le chat)
    # ============================================
    def process(self, message: str, history: List = None) -> str:
        """
        Méthode appelée par le chat pour traiter une demande.
        Extrait le nom du fichier du message et le lit.
        """
        # Chercher le nom du fichier dans le message
        match = re.search(r'([\w\-\.]+\.[a-zA-Z0-9]+)', message)
        
        if match:
            filename = match.group(1)
            file_path = os.path.join(self.upload_folder, filename)
            
            if os.path.exists(file_path):
                try:
                    content = self.read_file(file_path)
                    return f"📄 **Contenu de {filename}**\n\n{content}"
                except Exception as e:
                    return f"❌ Erreur lors de la lecture: {str(e)}"
            else:
                return f"❌ Fichier '{filename}' non trouvé dans {self.upload_folder}/"
        else:
            return "⚠️ Spécifiez un nom de fichier. Exemple: 'Lis le fichier test.txt'"

    # ============================================
    # MÉTHODES DE LECTURE PAR FORMAT
    # ============================================
    
    def read_txt(self, file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        return f"📄 {os.path.basename(file_path)}\n\n{content[:5000]}" + ("..." if len(content) > 5000 else "")

    def read_csv(self, file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            rows = list(reader)
        
        preview = ""
        for i, row in enumerate(rows[:10]):
            preview += f"  Ligne {i+1}: {', '.join(row)}\n"
        
        return f"""📊 CSV: {os.path.basename(file_path)}
📏 {len(rows)} lignes, {len(rows[0]) if rows else 0} colonnes

📋 Aperçu (10 premières lignes):
{preview}"""

    def read_json(self, file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        content = json.dumps(data, indent=2, ensure_ascii=False)
        if len(content) > 3000:
            content = content[:3000] + "\n... (tronqué)"
        return f"📄 JSON: {os.path.basename(file_path)}\n\n{content}"

    def read_pdf(self, file_path: str) -> str:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            num_pages = len(reader.pages)
            content = ""
            for i, page in enumerate(reader.pages[:5]):
                content += page.extract_text()
            result = f"📄 PDF: {os.path.basename(file_path)}\n📏 {num_pages} pages\n\n📋 Contenu:\n{content[:3000]}"
            return result + ("..." if len(content) > 3000 else "")
        except ImportError:
            return "⚠️ PyPDF2 non installé. Installez: pip install PyPDF2"
        except Exception as e:
            return f"❌ Erreur PDF: {str(e)}"

    def read_docx(self, file_path: str) -> str:
        try:
            from docx import Document
            doc = Document(file_path)
            content = "\n".join([p.text for p in doc.paragraphs])
            return f"📄 DOCX: {os.path.basename(file_path)}\n\n{content[:5000]}" + ("..." if len(content) > 5000 else "")
        except ImportError:
            return "⚠️ python-docx non installé. Installez: pip install python-docx"
        except Exception as e:
            return f"❌ Erreur DOCX: {str(e)}"

    def read_excel(self, file_path: str) -> str:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, data_only=True)
            sheet = wb.active
            content = f"📊 Excel: {os.path.basename(file_path)}\n📏 {sheet.max_row} lignes, {sheet.max_column} colonnes\n\n"
            for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if i > 20:
                    content += f"... ({sheet.max_row - 20} lignes supplémentaires)\n"
                    break
                content += f"  Ligne {i}: {', '.join([str(cell) if cell is not None else '' for cell in row])}\n"
            return content
        except ImportError:
            return "⚠️ openpyxl non installé. Installez: pip install openpyxl"
        except Exception as e:
            return f"❌ Erreur Excel: {str(e)}"

    def read_pptx(self, file_path: str) -> str:
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            content = f"📑 PowerPoint: {os.path.basename(file_path)}\n📏 {len(prs.slides)} diapositives\n\n"
            for i, slide in enumerate(prs.slides[:10]):
                content += f"\nDiapositive {i+1}:\n"
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        content += f"  {shape.text}\n"
            return content
        except ImportError:
            return "⚠️ python-pptx non installé. Installez: pip install python-pptx"
        except Exception as e:
            return f"❌ Erreur PowerPoint: {str(e)}"

    def read_image(self, file_path: str) -> str:
        try:
            from PIL import Image
            img = Image.open(file_path)
            return f"🖼️ Image: {os.path.basename(file_path)}\n📏 {img.width}x{img.height} pixels"
        except ImportError:
            return f"🖼️ Image: {os.path.basename(file_path)}\nℹ️ PIL non installé"
        except Exception as e:
            return f"❌ Erreur image: {str(e)}"

    def read_archive(self, file_path: str) -> str:
        try:
            import zipfile
            content = f"📦 Archive: {os.path.basename(file_path)}\n\n📋 Contenu:\n"
            with zipfile.ZipFile(file_path, 'r') as z:
                for info in z.infolist()[:20]:
                    content += f"  📄 {info.filename} ({info.file_size} octets)\n"
            return content
        except:
            return f"📦 Archive: {os.path.basename(file_path)}\n⚠️ Lecture non supportée (format compressé)"